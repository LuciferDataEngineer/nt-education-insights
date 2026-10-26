"""Shared parsers for suppressed cells, rates, and jurisdiction labels."""

from __future__ import annotations

import math
import re
from typing import Any

import pandas as pd

SUPPRESSED = {
    "",
    "-",
    "–",
    "—",
    ".",
    "..",
    "...",
    "np",
    "n/a",
    "na",
    "nan",
    "none",
    "*",
}


def is_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    try:
        if pd.isna(value):
            return True
    except TypeError:
        pass
    return False


def _text(value: Any) -> str:
    return str(value).strip().replace("\xa0", " ").replace("\u200b", "")


def parse_number(value: Any) -> float | None:
    """Parse an enrolment or count. Suppressed tokens become null, never zero."""
    if is_missing(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if math.isnan(number):
            return None
        return number
    text = _text(value).lower().replace(",", "").replace(" ", "")
    if text in SUPPRESSED or text.startswith("np"):
        return None
    text = text.rstrip("%")
    try:
        return float(text)
    except ValueError:
        return None


def parse_rate(value: Any) -> float | None:
    """Parse an attendance or share into a 0–1 proportion.

    NT workbooks mix '81.5%', raw percents (81.5) and proportions (0.815).
    """
    if is_missing(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if math.isnan(number):
            return None
        if number > 1.5:
            number = number / 100.0
        if 0 <= number <= 1:
            return number
        return None
    text = _text(value).lower().replace(",", "").replace(" ", "")
    if text in SUPPRESSED or text.startswith("np"):
        return None
    percent = text.endswith("%")
    if percent:
        text = text[:-1]
    try:
        number = float(text)
    except ValueError:
        return None
    if percent or number > 1.5:
        number = number / 100.0
    if 0 <= number <= 1:
        return number
    return None


def parse_percent_points(value: Any) -> float | None:
    """Return a 0–100 percentage, preserving ACARA-style percent fields."""
    rate = parse_rate(value)
    if rate is None:
        return None
    return rate * 100.0


def clean_school_name(value: Any) -> tuple[str, int]:
    """Strip footnote asterisks. The count flags RSAS/SEAM marks in early files."""
    raw = _text(value)
    match = re.search(r"\*+$", raw)
    markers = len(match.group(0)) if match else 0
    name = re.sub(r"\*+$", "", raw).strip()
    name = re.sub(r"\s+", " ", name)
    return name, markers


def normalise_name(value: str) -> str:
    text = value.lower().replace("&", " and ")
    text = text.replace("’", "'")
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def clean_abs_label(value: Any) -> str | None:
    if is_missing(value):
        return None
    text = _text(value)
    if text.startswith("(") or text.startswith("©"):
        return None
    text = re.sub(r"^[a-zA-Z]\s+", "", text).strip()
    text = text.rstrip(".").strip()
    return text or None


def canonical_state(value: Any) -> str | None:
    label = clean_abs_label(value)
    if label is None:
        return None
    key = label.lower().replace(".", "")
    mapping = {
        "nsw": "NSW",
        "new south wales": "NSW",
        "vic": "Vic",
        "victoria": "Vic",
        "qld": "Qld",
        "queensland": "Qld",
        "sa": "SA",
        "south australia": "SA",
        "wa": "WA",
        "western australia": "WA",
        "tas": "Tas",
        "tasmania": "Tas",
        "nt": "NT",
        "northern territory": "NT",
        "act": "ACT",
        "australian capital territory": "ACT",
        "aust": "Aust",
        "australia": "Aust",
    }
    return mapping.get(key)


def canonical_jurisdiction(value: Any) -> str | None:
    if is_missing(value):
        return None
    text = _text(value)
    text = re.sub(r"\s*\([^)]*\)", "", text).strip()
    return canonical_state(text)


def year_from_label(value: Any) -> int | None:
    if is_missing(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        year = int(value)
        if 1990 <= year <= 2100:
            return year
        return None
    match = re.search(r"(20\d{2}|19\d{2})", _text(value))
    if not match:
        return None
    return int(match.group(1))


def financial_year_label(value: Any) -> str | None:
    if is_missing(value):
        return None
    text = _text(value)
    match = re.search(r"(20\d{2})\s*[-–]\s*(\d{2})", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    year = year_from_label(text)
    if year:
        return str(year)
    return None


def write_parquet(frame: pd.DataFrame, path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_parquet(path, index=False)


def read_excel_rows(path) -> tuple[str, list[tuple]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet_name = workbook.sheetnames[0]
    sheet = workbook[sheet_name]
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return sheet_name, rows
